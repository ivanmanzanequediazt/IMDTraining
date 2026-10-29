"""Plantillas IMDTraining: una pestaña por semana con 4 tablas (una por sesión)."""
import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as GL

F = 'Arial'
AZUL = '1F4E79'      # títulos de sesión
AZUL2 = '2F6F9F'     # cabecera de columnas
VERDE = 'E3F4EA'     # filas alternas
VERDE2 = 'C6E9D3'    # celdas a rellenar (día, nombre, fecha)
CIELO = 'DCEBF7'     # bloque aeróbico
MENTA = 'D4F0DF'     # bloque activación
TXT = '1B2A38'
line = Side(style='thin', color='B7CCDD')
DIAS = 'Lunes,Martes,Miércoles,Jueves,Viernes,Sábado,Domingo'
BLOQUES = 'Activación,P. Principal,Aeróbico,Accesorios,Vuelta a la calma'
FILAS = 15           # ejercicios por sesión
N_SES = 4
fill = lambda c: PatternFill('solid', fgColor=c)


def build(path, pareja):
    pesos = ['Peso O', 'Peso S'] if pareja else ['Peso (kg)']
    cols = ['Bloque', 'Ejercicio', 'Series', 'Reps', 'Nota reps'] + pesos + ['RPE', 'Unilateral', 'Pliometría', 'Observaciones']
    widths = {'Bloque': 17, 'Ejercicio': 34, 'Series': 8, 'Reps': 8, 'Nota reps': 13, 'Peso (kg)': 10, 'Peso O': 9, 'Peso S': 9, 'RPE': 7,
              'Unilateral': 14, 'Pliometría': 10, 'Observaciones': 40}
    idx = {c: i + 1 for i, c in enumerate(cols)}
    centro = {'Series', 'Reps', 'RPE', 'Unilateral', 'Pliometría'} | set(pesos)
    wb = Workbook()

    # ---------- INSTRUCCIONES ----------
    ins = wb.active; ins.title = 'INSTRUCCIONES'
    ins.sheet_properties.tabColor = AZUL
    ins['A1'] = 'Cómo se rellena esta plantilla'; ins['A1'].font = Font(name=F, bold=True, size=15, color=AZUL)
    lines = [
        'Una pestaña por semana (Semana 1 … Semana 5) con 4 tablas, una por sesión. Si el cliente entrena menos días, deja vacías las que sobren.',
        'Celda verde «SEMANA DEL» de «Semana 1»: el LUNES de la primera semana del mes. Las demás semanas se calculan solas.',
        'En cada tabla: elige el Día (desplegable) y, si quieres, un Nombre (Pierna, Tracc, Empuje…). Sin nombre será «Entreno lunes».',
        'RPE pre y RPE post (arriba de cada tabla): para anotar lo que os diga el cliente. No pasan a la app; allí lo anota el cliente.',
        'Bloque: Activación / P. Principal / Aeróbico / Accesorios / Vuelta a la calma. Se pone al empezar cada bloque; las filas de debajo heredan.',
        'A la app pasan la Activación (lista para marcar, sin carga), la P. Principal y el Aeróbico. Accesorios y Vuelta a la calma se quedan en el Excel.',
        'Series y Reps: solo números. Lo que no sea número (6-8, EA, por lado, 3x4") va en «Nota reps».',
        'Biseries: dos filas seguidas (una por ejercicio) o «Ejercicio A + Ejercicio B» en una fila.',
        'Unilateral / Pliometría: elige «Sí» si lo es. En pliometría, Reps = contactos.',
        ('Peso O / Peso S: un peso por cliente. Cambia la letra del título por la inicial de cada uno (Peso A, Peso M…) '
         'y usa esas mismas letras en la pestaña «fuentes» de Google (por ejemplo «O: apodo de uno, S: apodo del otro»).') if pareja else
        'Peso (kg): solo el número (17,5). Si hace falta texto («por lado», «poner lastre»), va en Observaciones.',
        'Aeróbico (al final de la sesión): Bloque «Aeróbico», Ejercicio = Carrera, Bicicleta…, y en Observaciones el tiempo y la zona: «20 min Z2» o, por intervalos, «10 Z2, 3 Z4, 10 Z2». Una fila por sesión.',
        'No cambies los títulos de las tablas ni la celda «SEMANA DEL». Puedes añadir más pestañas (ANUAL, RM…): se ignoran.',
        'Para el mes siguiente: Archivo > Guardar una copia, cambia la fecha verde y borra el contenido de las tablas.',
    ]
    for i, t in enumerate(lines):
        c = ins.cell(3 + i, 1, f'{i + 1}. {t}'); c.font = Font(name=F, size=10, color=TXT); c.alignment = Alignment(wrap_text=True, vertical='top')
        ins.merge_cells(start_row=3 + i, start_column=1, end_row=3 + i, end_column=len(cols))
        ins.row_dimensions[3 + i].height = 28
        if i % 2:
            for k in range(1, len(cols) + 1): ins.cell(3 + i, k).fill = fill(VERDE)
    for c in cols: ins.column_dimensions[GL(idx[c])].width = widths[c]
    r0 = 4 + len(lines)
    ins.cell(r0, 1, 'Ejemplo (esta pestaña no se importa):').font = Font(name=F, bold=True, size=11, color=AZUL)
    ej = [
        {'Bloque': 'Activación', 'Ejercicio': 'Puente glúteo goma', 'Series': 2, 'Reps': 10},
        {'Bloque': 'P. Principal', 'Ejercicio': 'Sentadilla goblet', 'Series': 3, 'Reps': 8, 'Nota reps': '8-10', 'RPE': 7,
         **({'Peso O': 24, 'Peso S': 12} if pareja else {'Peso (kg)': 24})},
        {'Ejercicio': 'Zancada atrás', 'Series': 3, 'Reps': 8, 'Nota reps': 'por lado', 'Unilateral': 'Sí',
         **({'Peso O': 10, 'Peso S': 6} if pareja else {'Peso (kg)': 10})},
        {'Ejercicio': 'Drop jump', 'Series': 3, 'Reps': 6, 'Pliometría': 'Sí'},
        {'Bloque': 'Aeróbico', 'Ejercicio': 'Bicicleta', 'Observaciones': '15 min Z2'},
    ]
    tabla(ins, r0 + 1, 1, cols, idx, centro, pesos, len(ej), titulo=False, ejemplo=ej)

    # ---------- Semanas ----------
    for w in range(1, 6):
        ws = wb.create_sheet(f'Semana {w}')
        ws.sheet_properties.tabColor = '2E8B6E' if w % 2 else '2F6F9F'
        ws.sheet_view.showGridLines = False
        ws['A1'] = 'SEMANA DEL'; ws['A1'].font = Font(name=F, bold=True, color=AZUL, size=12)
        if w == 1:
            ws['B1'] = datetime.date(2026, 11, 2)
            ws['B1'].comment = Comment('Escribe aquí el LUNES de la primera semana del mes (las demás semanas se calculan solas).', 'IMDTraining')
        else:
            ws['B1'] = f"='Semana {w - 1}'!B1+7"
        ws['B1'].number_format = 'DD/MM/YYYY'; ws['B1'].font = Font(name=F, bold=True, size=13, color=TXT)
        ws['B1'].fill = fill(VERDE2 if w == 1 else 'FFFFFF'); ws['B1'].alignment = Alignment(horizontal='center')
        ws['C1'] = '="al "&TEXT(B1+6,"dd/mm")'; ws['C1'].font = Font(name=F, size=12, color='4A6A85')
        ws['E1'] = 'Cliente'; ws['E1'].font = Font(name=F, color='4A6A85'); ws['E1'].alignment = Alignment(horizontal='right')
        ws['F1'].fill = fill(VERDE); ws['F1'].font = Font(name=F, color=TXT)
        ws.merge_cells('F1:H1')
        for c in cols: ws.column_dimensions[GL(idx[c])].width = widths[c]
        r = 3
        for k in range(1, N_SES + 1):
            tabla(ws, r, k, cols, idx, centro, pesos, FILAS)
            r += FILAS + 4
        ws.freeze_panes = 'A2'
        ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 1
    wb.save(path)


def tabla(ws, r, k, cols, idx, centro, pesos, nfilas, titulo=True, ejemplo=None):
    """Fila r: título de la sesión (Día, Nombre, RPE pre/post). r+1: cabecera. r+2…: ejercicios."""
    last = len(cols)
    L = GL
    if titulo:
        for c in range(1, last + 1):
            ws.cell(r, c).fill = fill(AZUL)
        t = ws.cell(r, 1, f'SESIÓN {k}'); t.font = Font(name=F, bold=True, size=12, color='FFFFFF')
        u, pl, ob = idx['Unilateral'], idx['Pliometría'], idx['Observaciones']
        meta = [(2, 'Día:', 3), (5, 'Nombre:', 6), (u, 'RPE pre / post:', pl), (None, None, ob)]
        for lc, lab, vc in meta:
            if lc:
                a = ws.cell(r, lc, lab); a.font = Font(name=F, bold=True, color='FFFFFF', size=9 if len(lab) > 8 else 10); a.alignment = Alignment(horizontal='right', vertical='center')
            v = ws.cell(r, vc); v.fill = fill(VERDE2); v.font = Font(name=F, bold=True, color=TXT); v.alignment = Alignment(horizontal='center', vertical='center')
            v.border = Border(left=line, right=line, top=line, bottom=line)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        if u - 1 > 6: ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=u - 1)
        ws.row_dimensions[r].height = 22
        dv = DataValidation(type='list', formula1=f'"{DIAS}"', allow_blank=True, showErrorMessage=True, errorTitle='Día', error='Elige un día del desplegable.')
        ws.add_data_validation(dv); dv.add(f'C{r}')
        dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='10', allow_blank=True, showErrorMessage=True, error='Un número de 0 a 10.')
        ws.add_data_validation(dv); dv.add(f'{L(pl)}{r}'); dv.add(f'{L(ob)}{r}')
        h = r + 1
    else:
        h = r
    for i, c in enumerate(cols):
        cell = ws.cell(h, i + 1, c)
        cell.font = Font(name=F, bold=True, color='FFFFFF', size=10); cell.fill = fill(AZUL2)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[h].height = 20
    a, b = h + 1, h + nfilas
    for rr in range(a, b + 1):
        for c in cols:
            cell = ws.cell(rr, idx[c]); cell.font = Font(name=F, size=10, color=TXT)
            cell.border = Border(bottom=line, left=line if c == cols[0] else None, right=line if c == cols[-1] else None)
            if (rr - a) % 2: cell.fill = fill(VERDE)
            if c in centro: cell.alignment = Alignment(horizontal='center')
        if ejemplo and rr - a < len(ejemplo):
            for c, v in ejemplo[rr - a].items(): ws.cell(rr, idx[c], v)
    if ejemplo:
        return
    rng = lambda c: f'{L(idx[c])}{a}:{L(idx[c])}{b}'

    def dv(c, **kw):
        v = DataValidation(allow_blank=True, showErrorMessage=True, errorStyle='stop', **kw); ws.add_data_validation(v); v.add(rng(c))
    dv('Bloque', type='list', formula1=f'"{BLOQUES}"', errorTitle='Bloque', error='Elige un bloque del desplegable.')
    dv('Series', type='whole', operator='between', formula1='1', formula2='20', errorTitle='Series', error='Solo un número de 1 a 20.')
    dv('Reps', type='whole', operator='between', formula1='1', formula2='300', errorTitle='Reps', error='Solo números. El texto (6-8, EA, por lado) va en «Nota reps».')
    for pz in pesos:
        dv(pz, type='decimal', operator='between', formula1='0', formula2='1000', errorTitle='Peso', error='Solo el número de kilos (por ejemplo 17,5). El texto va en Observaciones.')
    dv('RPE', type='decimal', operator='between', formula1='0', formula2='10', errorTitle='RPE', error='Un número de 0 a 10.')
    dv('Unilateral', type='list', formula1='"Sí"', errorTitle='Unilateral', error='Elige «Sí» o déjalo vacío.')
    dv('Pliometría', type='list', formula1='"Sí"', errorTitle='Pliometría', error='Elige «Sí» o déjalo vacío.')
    bc = L(idx['Bloque'])
    full = f'{bc}{a}:{bc}{b}'
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${bc}{a}="Activación"'], fill=fill(MENTA), font=Font(bold=True, color='1E6B45')))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${bc}{a}="P. Principal"'], fill=fill('D6E6F5'), font=Font(bold=True, color=AZUL)))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${bc}{a}="Aeróbico"'], fill=fill(CIELO), font=Font(bold=True, color='1F5C8A')))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'OR(${bc}{a}="Accesorios",${bc}{a}="Vuelta a la calma")'], font=Font(italic=True, color='6B7C8C')))


build(sys.argv[1] + '/PLANTILLA_CLIENTE.xlsx', False)
build(sys.argv[1] + '/PLANTILLA_PAREJA.xlsx', True)
print('ok')
